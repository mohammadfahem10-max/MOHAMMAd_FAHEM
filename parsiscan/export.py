"""Write results as .txt, .json and .docx (right-to-left Word document)."""
from __future__ import annotations

import os

from .pipeline import DocumentOut, UNKNOWN_MARK


def write_txt(doc: DocumentOut, path: str) -> None:
    parts = []
    for p in doc.pages:
        header = f"===== صفحه {p.index + 1} ====="
        if p.source == "handwritten":
            parts.append(header + "\n[دست‌نویس: با روش تطبیق الگو قابل خواندن نیست؛ نیاز به بازبینی انسانی]\n")
        elif p.source == "empty":
            parts.append(header + "\n[صفحه‌ی خالی]\n")
        else:
            parts.append(header + "\n" + p.text + "\n")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(parts))


def write_docx(doc: DocumentOut, path: str) -> bool:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import RGBColor
    except ImportError:
        return False

    d = Document()
    style = d.styles["Normal"]
    style.font.name = "Tahoma"
    style.element.rPr.rFonts.set(qn("w:cs"), "Tahoma")

    def rtl_paragraph(text: str, highlight_unknown: bool = True, note: bool = False):
        para = d.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        ppr = para._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
        if note:
            run = para.add_run(text)
            run.italic = True
            run.font.color.rgb = RGBColor(0x99, 0x33, 0x00)
            return
        # colour the unknown marks so they are easy to find in Word
        chunks = text.split(UNKNOWN_MARK)
        for k, chunk in enumerate(chunks):
            if chunk:
                para.add_run(chunk)
            if k < len(chunks) - 1:
                run = para.add_run(UNKNOWN_MARK)
                run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
                run.bold = True

    for p in doc.pages:
        d.add_heading(f"صفحه {p.index + 1}", level=2)
        if p.source == "handwritten":
            rtl_paragraph("این صفحه دست‌نویس تشخیص داده شد و با تطبیق الگو خوانده نمی‌شود.", note=True)
            continue
        if p.source == "empty":
            rtl_paragraph("صفحه‌ی خالی", note=True)
            continue
        for line in p.text.split("\n"):
            rtl_paragraph(line)
        bad = [n for n in p.numbers if not n["valid"]]
        if bad:
            rtl_paragraph("اعداد مشکوک: " + "، ".join(f"{n['text']} ({n['note']})" for n in bad), note=True)
    d.save(path)
    return True


def write_outputs(doc: DocumentOut, out_dir: str) -> dict:
    os.makedirs(out_dir, exist_ok=True)
    txt = os.path.join(out_dir, "result.txt")
    js = os.path.join(out_dir, "result.json")
    dx = os.path.join(out_dir, "result.docx")
    write_txt(doc, txt)
    with open(js, "w", encoding="utf-8") as fh:
        fh.write(doc.to_json())
    ok = write_docx(doc, dx)
    return {"txt": txt, "json": js, "docx": dx if ok else None}
